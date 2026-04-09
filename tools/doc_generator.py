import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from datetime import datetime
from typing import Dict, Any, List, Optional
from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from config import settings


class DocumentGenerator:
    
    def __init__(self):
        self.template_path = settings.TEMPLATES_DIR / "report_template.docx"
    
    def create_document(
        self,
        input_data: Dict[str, Any],
        parsed_data: Dict[str, Any],
        solutions: List[Dict[str, Any]],
        output_path: Optional[str] = None
    ) -> str:
        doc = Document()
        
        self._set_document_style(doc)
        
        title = doc.add_heading("制造偏差处理方案报告", 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc.add_paragraph("")
        
        info_para = doc.add_paragraph()
        info_para.add_run("报告编号: ").bold = True
        report_id = f"AI-DTS-{datetime.now().strftime('%Y%m%d%H%M%S')}"
        info_para.add_run(report_id)
        
        info_para = doc.add_paragraph()
        info_para.add_run("生成时间: ").bold = True
        info_para.add_run(datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
        
        doc.add_paragraph("")
        doc.add_heading("一、输入数据", level=1)
        
        table = doc.add_table(rows=1, cols=2)
        table.style = 'Table Grid'
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        header_cells = table.rows[0].cells
        header_cells[0].text = "数据项"
        header_cells[1].text = "数值"
        self._set_cell_shading(header_cells[0])
        self._set_cell_shading(header_cells[1])
        
        data_items = [
            ("监理公司", parsed_data.get("supervisor", "N/A")),
            ("铰点型号", parsed_data.get("hinge_type", "N/A")),
        ]
        
        if "deviation_1_4" in parsed_data:
            data_items.append(("序号1-4偏心量", f"{parsed_data['deviation_1_4']} mm"))
            data_items.append(("序号1-4偏心量区间", parsed_data.get("deviation_1_4_range", "N/A")))
        
        if "deviation_5_8" in parsed_data:
            data_items.append(("序号5-8偏心量", f"{parsed_data['deviation_5_8']} mm"))
            data_items.append(("序号5-8偏心量区间", parsed_data.get("deviation_5_8_range", "N/A")))
        
        for item_name, item_value in data_items:
            row = table.add_row()
            row.cells[0].text = item_name
            row.cells[1].text = str(item_value) if item_value else "N/A"
        
        doc.add_paragraph("")
        doc.add_heading("二、计算过程", level=1)
        
        calculations = parsed_data.get("calculations", {})
        for calc_name, calc_data in calculations.items():
            para = doc.add_paragraph()
            para.add_run(f"{calc_name}: ").bold = True
            para.add_run(calc_data.get("formula", "N/A"))
        
        doc.add_paragraph("")
        doc.add_heading("三、处理方案", level=1)
        
        for i, solution in enumerate(solutions, 1):
            doc.add_heading(f"方案 {i}", level=2)
            
            table = doc.add_table(rows=4, cols=2)
            table.style = 'Table Grid'
            
            rows_data = [
                ("规则ID", solution.get("rule_id", "N/A")),
                ("数据来源", solution.get("source", "N/A")),
                ("置信度", solution.get("confidence", "N/A")),
                ("处理方案", solution.get("solution", "N/A")),
            ]
            
            for row_idx, (label, value) in enumerate(rows_data):
                table.rows[row_idx].cells[0].text = label
                table.rows[row_idx].cells[1].text = str(value) if value else "N/A"
                self._set_cell_shading(table.rows[row_idx].cells[0])
            
            doc.add_paragraph("")
        
        doc.add_heading("四、备注", level=1)
        note = doc.add_paragraph()
        note.add_run("本报告由AI-DTS制造偏差处理方案自动生成系统生成，仅供参考。")
        note.add_run("\n如有疑问，请联系相关工程师进行人工审核。")
        
        if output_path is None:
            output_path = str(settings.OUTPUT_DIR / f"report_{report_id}.docx")
        
        doc.save(output_path)
        
        return output_path
    
    def _set_document_style(self, doc):
        style = doc.styles['Normal']
        style.font.name = 'SimSun'
        style._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')
        style.font.size = Pt(12)
    
    def _set_cell_shading(self, cell):
        shading = OxmlElement('w:shd')
        shading.set(qn('w:fill'), 'E6E6E6')
        cell._tc.get_or_add_tcPr().append(shading)


def generate_report(
    input_data: Dict[str, Any],
    parsed_data: Dict[str, Any],
    solutions: List[Dict[str, Any]]
) -> str:
    generator = DocumentGenerator()
    return generator.create_document(input_data, parsed_data, solutions)


if __name__ == "__main__":
    test_parsed = {
        "supervisor": "GBB",
        "hinge_type": "1#铰点",
        "deviation_1_4": 5.0,
        "deviation_1_4_range": "<=5mm",
        "calculations": {
            "序号1-4": {
                "formula": "C = sqrt(3.0^2 + 4.0^2) = 5.0mm"
            }
        }
    }
    
    test_solutions = [
        {
            "rule_id": "序号1-4规则组_规则1",
            "source": "知识图谱",
            "confidence": "高",
            "solution": "处理结果：按图纸尺寸划线。"
        }
    ]
    
    output = generate_report({}, test_parsed, test_solutions)
    print(f"报告已生成: {output}")
